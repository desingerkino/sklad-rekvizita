# -*- coding: utf-8 -*-
"""
Экспорт данных в Excel (.xlsx).

Файлы делаются библиотекой openpyxl — это настоящие книги Excel со шрифтами,
шириной колонок, закреплённой шапкой, автофильтром и числовыми форматами, а не
переименованный CSV. Открываются и в Excel, и в LibreOffice, и в «Google
Таблицах».

PDF никуда не девается: Excel добавлен как второй формат для случаев, когда с
выгрузкой нужно работать дальше — считать, фильтровать, сводить.
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional, Sequence

from ..db.database import get_db
from ..errors import AppError
from ..utils import fmt_date, fmt_dt
from .audit_service import audit_service
from .auth_service import auth_service

log = logging.getLogger(__name__)

MONEY_FORMAT = "#,##0.00"
HEADER_FILL = "F1F5F9"


def _openpyxl():
    try:
        import openpyxl                                    # noqa: F401
        return openpyxl
    except ImportError as exc:                             # pragma: no cover
        raise AppError(
            "Библиотека openpyxl не установлена — выгрузка в Excel недоступна. "
            "Переустановите программу или сообщите администратору.") from exc


class ExcelService:
    # ----------------------------------------------------------- оформление
    def _sheet(self, workbook, title: str, headers: Sequence[str], widths: Sequence[int]):
        from openpyxl.styles import Alignment, Font, PatternFill

        sheet = workbook.create_sheet(title=title[:31])
        sheet.append(list(headers))
        header_font = Font(bold=True)
        fill = PatternFill("solid", fgColor=HEADER_FILL)
        for index, cell in enumerate(sheet[1], start=1):
            cell.font = header_font
            cell.fill = fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if index <= len(widths):
                sheet.column_dimensions[cell.column_letter].width = widths[index - 1]
        sheet.freeze_panes = "A2"          # шапка не уезжает при прокрутке
        sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).column_letter}1"
        return sheet

    @staticmethod
    def _money_column(sheet, columns: Iterable[int]) -> None:
        for column in columns:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column).number_format = MONEY_FORMAT

    @staticmethod
    def _finish(workbook, target: Path, what: str, rows: int) -> Path:
        target = Path(target)
        target.parent.mkdir(parents=True, exist_ok=True)
        if "Sheet" in workbook.sheetnames:          # пустой лист по умолчанию
            del workbook["Sheet"]
        workbook.save(str(target))
        audit_service.log("Экспорт данных", details=f"{what} → Excel, строк: {rows}")
        log.info("Экспорт %s в %s (%s строк)", what, target, rows)
        return target

    @staticmethod
    def _title_row(sheet, text: str) -> None:
        """Подпись под таблицей: когда и кем выгружено."""
        user = auth_service.current_user
        sheet.append([])
        sheet.append([f"{text}. Выгружено {datetime.now():%d.%m.%Y %H:%M}"
                      + (f", пользователь: {user.full_name}" if user else "")])

    # ---------------------------------------------------------------- склад
    def export_inventory(self, target: Path, item_ids: Optional[Sequence[int]] = None,
                         include_archived: bool = False) -> Path:
        """
        Выгрузка каталога: остатки, стоимость и текущая аренда по каждой позиции.
        Если передан item_ids — выгружаются только выбранные строки.
        """
        auth_service.require_login()
        openpyxl = _openpyxl()
        from .inventory_service import inventory_service
        from .packing_service import packing_service

        db = get_db()
        sql = ("SELECT i.id, i.inv_number, i.barcode, i.name, i.description, i.status, "
               "       i.quantity_total, i.rental_price, i.valuation, i.is_archived, "
               "       ifnull(c.name,'') AS category "
               "FROM inventory_items i LEFT JOIN categories c ON c.id=i.category_id ")
        conditions, params = [], []
        if not include_archived:
            conditions.append("i.is_archived=0")
        if item_ids:
            placeholders = ",".join("?" * len(item_ids))
            conditions.append(f"i.id IN ({placeholders})")
            params += list(item_ids)
        if conditions:
            sql += "WHERE " + " AND ".join(conditions) + " "
        sql += "ORDER BY i.inv_number"

        workbook = openpyxl.Workbook()
        sheet = self._sheet(
            workbook, "Склад",
            ["Инв. №", "Штрихкод", "Наименование", "Категория", "Статус", "Всего",
             "Доступно", "В аренде", "В ремонте", "Цена аренды", "Оценка",
             "Место упаковки", "Текущая аренда", "Описание"],
            [12, 14, 38, 24, 14, 8, 10, 10, 11, 13, 13, 28, 34, 45])

        rows = 0
        for row in db.query(sql, params):
            rented = inventory_service.reserved_now(row["id"])
            repair = inventory_service.in_repair(row["id"])
            available = max(0, int(row["quantity_total"]) - rented - repair)
            rental = db.query_one(
                "SELECT r.number, ifnull(cl.full_name,'') AS client "
                "FROM rental_items ri JOIN rentals r ON r.id=ri.rental_id "
                "LEFT JOIN clients cl ON cl.id=r.client_id "
                "WHERE ri.item_id=? AND ri.quantity > ri.returned_qty "
                "  AND r.status IN ('active','overdue') ORDER BY ri.id LIMIT 1", (row["id"],))
            sheet.append([
                row["inv_number"], row["barcode"], row["name"], row["category"],
                "в архиве" if row["is_archived"] else row["status"],
                int(row["quantity_total"]), available, rented, repair,
                float(row["rental_price"] or 0), float(row["valuation"] or 0),
                packing_service.where_is_text(row["id"]),
                f"{rental['number']} · {rental['client']}" if rental else "",
                row["description"],
            ])
            rows += 1

        self._money_column(sheet, (10, 11))
        self._title_row(sheet, "Выгрузка склада")
        return self._finish(workbook, target, "Склад", rows)

    # --------------------------------------------------------------- аренды
    def export_rentals(self, target: Path, rental_ids: Optional[Sequence[int]] = None,
                       date_from: str = "", date_to: str = "") -> Path:
        """
        Выгрузка аренд: два листа — сами аренды и их позиции.

        Позиции отдельным листом, потому что в одной аренде их могут быть
        десятки: в одной таблице они либо потерялись бы, либо размножили
        строку аренды.
        """
        auth_service.require_login()
        openpyxl = _openpyxl()
        db = get_db()

        sql = ("SELECT r.*, ifnull(c.full_name,'') AS client, "
               "       ifnull(p.name,'') AS project, ifnull(pr.name,'') AS production, "
               "       (SELECT a.user_name FROM audit_log a WHERE a.entity='rental' "
               "         AND a.entity_id=r.id AND a.action LIKE 'Создана%' "
               "         ORDER BY a.id LIMIT 1) AS user_name "
               "FROM rentals r "
               "LEFT JOIN clients c ON c.id=r.client_id "
               "LEFT JOIN projects p ON p.id=r.project_id "
               "LEFT JOIN productions pr ON pr.id=r.production_id ")
        conditions, params = [], []
        if rental_ids:
            conditions.append(f"r.id IN ({','.join('?' * len(rental_ids))})")
            params += list(rental_ids)
        if date_from:
            conditions.append("r.date_start >= ?")
            params.append(date_from)
        if date_to:
            conditions.append("r.date_start <= ?")
            params.append(date_to)
        if conditions:
            sql += "WHERE " + " AND ".join(conditions) + " "
        sql += "ORDER BY r.date_start DESC, r.id DESC"

        workbook = openpyxl.Workbook()
        head = self._sheet(
            workbook, "Аренды",
            ["Номер", "Статус", "Клиент", "Продакшен", "Проект", "Выдана",
             "План. возврат", "Факт. возврат", "Позиций", "Стоимость", "Скидка, %",
             "Скидка", "Итого", "Залог", "Оплата", "Ответственный", "Примечание"],
            [16, 12, 30, 28, 28, 13, 14, 14, 9, 13, 11, 13, 13, 12, 12, 24, 30])
        lines = self._sheet(
            workbook, "Позиции аренд",
            ["Аренда", "Клиент", "Инв. №", "Наименование", "Количество", "Возвращено",
             "Цена", "Сумма", "Выдана", "Период с", "Период по", "Пролонгаций",
             "Примечание"],
            [16, 28, 12, 38, 12, 12, 13, 13, 13, 13, 13, 12, 28])

        rows = 0
        for rental in db.query(sql, params):
            items = db.query(
                "SELECT * FROM rental_items WHERE rental_id=? ORDER BY id", (rental["id"],))
            returned = db.query_one(
                "SELECT MAX(returned_at) v FROM rental_returns rr "
                "JOIN rental_items ri ON ri.id=rr.rental_item_id WHERE ri.rental_id=?",
                (rental["id"],))
            fully_returned = all(
                int(i["returned_qty"]) >= int(i["quantity"]) for i in items) if items else False
            head.append([
                rental["number"], rental["status"], rental["client"], rental["production"],
                rental["project"], fmt_date(rental["date_start"]),
                fmt_date(rental["date_end"]),
                fmt_dt(returned["v"]) if returned and returned["v"] and fully_returned else "",
                len(items),
                float(rental["total_before_discount"] or 0),
                float(rental["discount_percent"] or 0),
                float(rental["discount_amount"] or 0),
                float(rental["total_final"] or 0),
                float(rental["deposit"] or 0),
                rental["payment_method"], rental["user_name"] or "", rental["note"],
            ])
            rows += 1
            for item in items:
                lines.append([
                    rental["number"], rental["client"], item["inv_number_snap"],
                    item["name_snap"], int(item["quantity"]), int(item["returned_qty"]),
                    float(item["price"] or 0),
                    float(item["price"] or 0) * int(item["quantity"]),
                    fmt_date(item["date_start"]) if item["date_start"] else "",
                    fmt_date(item["period_start"]) if item["period_start"] else "",
                    fmt_date(item["period_end"]) if item["period_end"] else "",
                    int(item["prolongations"] or 0), item["note"],
                ])

        self._money_column(head, (10, 12, 13, 14))
        self._money_column(lines, (7, 8))
        self._title_row(head, "Выгрузка аренд")
        return self._finish(workbook, target, "Аренды", rows)

    # ------------------------------------------------------- инвентаризация
    def export_audit(self, target: Path, audit_id: int) -> Path:
        """Выгрузка инвентаризации: итоги и построчная сверка."""
        auth_service.require_login()
        openpyxl = _openpyxl()
        from .inventory_audit_service import inventory_audit_service

        audit = inventory_audit_service.get(audit_id)
        summary = inventory_audit_service.summary(audit_id)
        lines = inventory_audit_service.lines(audit_id)

        workbook = openpyxl.Workbook()
        totals = self._sheet(workbook, "Итоги", ["Показатель", "Значение"], [42, 22])
        for caption, value in (
                ("Инвентаризация", audit["number"]),
                ("Название", audit["title"] or "—"),
                ("Состояние", audit["status_label"]),
                ("Начата", fmt_dt(audit["started_at"])),
                ("Завершена", fmt_dt(audit["finished_at"]) if audit["finished_at"] else "—"),
                ("Ответственный", audit["responsible"] or audit.get("user_name") or "—"),
                ("Числится на складе, шт.", summary["expected_units"]),
                ("Найдено, шт.", summary["found_units"]),
                ("Не найдено (недостача), шт.", summary["missing_units"]),
                ("Излишки, шт.", summary["extra_units"]),
                ("Позиций совпало", summary["positions_ok"]),
                ("Позиций с недостачей", summary["positions_missing"]
                 + summary["positions_partial"]),
                ("Процент соответствия", summary["match_percent"]),
                ("Справочно: у клиентов, шт.", summary["rented_units"]),
                ("Справочно: в ремонте, шт.", summary["repair_units"])):
            totals.append([caption, value])

        detail = self._sheet(
            workbook, "Сверка",
            ["Инв. №", "Наименование", "Категория", "Числится", "Найдено",
             "Расхождение", "У клиентов", "В ремонте", "Состояние", "Комментарий"],
            [12, 38, 24, 11, 11, 13, 12, 12, 16, 30])
        states = {"ok": "совпало", "missing": "не найдено", "partial": "недостача",
                  "extra": "излишек"}
        for line in lines:
            detail.append([
                line["inv_number"], line["name_snap"], line["category_snap"],
                int(line["expected_qty"]), int(line["found_qty"]), int(line["diff"]),
                int(line["rented_qty"]), int(line["repair_qty"]),
                states.get(line["state"], line["state"]), line["comment"],
            ])

        self._title_row(detail, f"Инвентаризация {audit['number']}")
        return self._finish(workbook, target, f"Инвентаризация {audit['number']}", len(lines))

    # ---------------------------------------------------------- комплектация
    def export_packing(self, target: Path, set_id: int) -> Path:
        """Выгрузка комплекта: коробки и их содержимое."""
        auth_service.require_login()
        openpyxl = _openpyxl()
        from .packing_service import packing_service

        data = packing_service.get(set_id)
        boxes = packing_service.contents(set_id)

        workbook = openpyxl.Workbook()
        sheet = self._sheet(
            workbook, "Комплектация",
            ["Коробка", "Инв. №", "Наименование", "Количество", "Примечание"],
            [24, 12, 40, 12, 30])
        rows = 0
        for box in boxes:
            if not box["items"]:
                sheet.append([box["label"], "", "(пусто)", 0, box["note"]])
                continue
            for item in box["items"]:
                sheet.append([box["label"], item["inv_number"], item["name_snap"],
                              int(item["quantity"]), item["note"]])
                rows += 1

        self._title_row(
            sheet,
            f"Комплект {data['number']} · {data['vehicle_name'] or 'машина не указана'} · "
            f"{fmt_date(data['date_ship']) if data['date_ship'] else 'дата не указана'}")
        return self._finish(workbook, target, f"Комплектация {data['number']}", rows)


excel_service = ExcelService()
